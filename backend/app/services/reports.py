import io
from datetime import datetime, timezone

LEVEL_ES = {"seguro_provisional": "Seguro provisional", "bajo": "Bajo", "medio": "Medio", "alto": "Alto", "critico": "Crítico", "desconocido": "Desconocido"}


def _rows(data):
    return [[a["created_at"][:19].replace("T", " "), a["domain"], LEVEL_ES.get(a["level"], a["level"]), a["score"], a["status"]] for a in data["alerts"]]


HEAD = ["Fecha (UTC)", "Dominio", "Nivel", "Puntaje", "Estado"]


def to_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["BankShield AI - Reporte", data["generated_at"]])
    ws.append([])
    ws.append(["Total de URL analizadas", data["total"]])
    ws.append(["Alertas abiertas", data["open_alerts"]])
    ws.append([])
    ws.append(["Nivel", "Cantidad"])
    for k, v in data["by_level"].items():
        ws.append([LEVEL_ES.get(k, k), v])
    for c in ws[6]:
        c.font = Font(bold=True)
    d = wb.create_sheet("Alertas")
    d.append(HEAD)
    for c in d[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="12304A")
    for r in _rows(data):
        d.append(r)
    for col, w in zip("ABCDE", (20, 40, 20, 10, 18)):
        d.column_dimensions[col].width = w
    ws.column_dimensions["A"].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(data: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="BankShield AI - Reporte")
    st = getSampleStyleSheet()
    els = [Paragraph("BankShield AI - Reporte ejecutivo", st["Title"]),
           Paragraph(f"Generado: {data['generated_at']}", st["Normal"]), Spacer(1, 12),
           Paragraph(f"URL analizadas: {data['total']} &nbsp;|&nbsp; Alertas abiertas: {data['open_alerts']}", st["Normal"]), Spacer(1, 12)]
    t = Table([["Nivel", "Cantidad"]] + [[LEVEL_ES.get(k, k), v] for k, v in data["by_level"].items()])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.grey)]))
    els += [t, Spacer(1, 16), Paragraph("Últimas alertas", st["Heading2"])]
    rows = [HEAD] + [[r[0], r[1][:32], r[2], r[3], r[4]] for r in _rows(data)[:40]]
    t2 = Table(rows, repeatRows=1)
    t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304A")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    els.append(t2)
    els += [Spacer(1, 12), Paragraph("Nota: una disminución de alertas no implica por sí sola mayor seguridad; puede indicar menor capacidad de detección.", st["Italic"])]
    doc.build(els)
    return buf.getvalue()


def to_docx(data: dict) -> bytes:
    from docx import Document
    d = Document()
    d.add_heading("BankShield AI - Informe técnico", 0)
    d.add_paragraph(f"Generado: {data['generated_at']}")
    d.add_paragraph(f"URL analizadas: {data['total']}. Alertas abiertas: {data['open_alerts']}.")
    d.add_heading("Distribución por nivel", 1)
    t = d.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Nivel", "Cantidad"
    for k, v in data["by_level"].items():
        c = t.add_row().cells
        c[0].text, c[1].text = LEVEL_ES.get(k, k), str(v)
    d.add_heading("Alertas recientes", 1)
    t2 = d.add_table(rows=1, cols=len(HEAD))
    t2.style = "Light Grid Accent 1"
    for i, h in enumerate(HEAD):
        t2.rows[0].cells[i].text = h
    for r in _rows(data)[:50]:
        cells = t2.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


BUILDERS = {
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (to_pdf, "application/pdf"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
}
